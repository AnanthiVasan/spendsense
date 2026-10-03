"""Build a plain-English Spendsense overview deck."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = "docs/Spendsense-overview.pptx"

INK = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF8, 0xFA, 0xFC)
EMERALD = RGBColor(0x05, 0x96, 0x69)
EMERALD_SOFT = RGBColor(0xEC, 0xFD, 0xF5)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE2, 0xE8, 0xF0)

W = Inches(13.333)
H = Inches(7.5)


def set_run(run, text, size, bold, color, font="Calibri"):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, text, left, top, width, height, size, bold, color, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    p = frame.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size, bold, color)
    return box


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def card(slide, left, top, width, height):
    shape = rect(slide, left, top, width, height, CARD)
    shape.line.color.rgb = LINE
    shape.line.width = Pt(1)
    return shape


def new_slide(prs, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, INK if dark else PAPER)
    return slide


def header(slide, kicker, title):
    rect(slide, 0, 0, Inches(0.18), H, EMERALD)
    add_text(slide, kicker.upper(), Inches(0.7), Inches(0.38), Inches(10), Inches(0.32), 14, True, EMERALD)
    add_text(slide, title, Inches(0.7), Inches(0.68), Inches(12), Inches(1.15), 34, True, INK)


def bullets(slide, items, top=Inches(1.8)):
    box = slide.shapes.add_textbox(Inches(0.85), top, Inches(11.6), Inches(5.1))
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.level = 0
        p.space_after = Pt(16)
        set_run(p.add_run(), item, 24, False, SLATE)
    return box


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "Spendsense"
    prs.core_properties.subject = "A plain-English guide to what Spendsense does"

    # 1 Title
    s = new_slide(prs, dark=True)
    rect(s, 0, 0, Inches(0.22), H, EMERALD)
    add_text(s, "PERSONAL FINANCE", Inches(0.85), Inches(1.85), Inches(10), Inches(0.35), 16, True, EMERALD)
    add_text(s, "Spendsense", Inches(0.85), Inches(2.25), Inches(11), Inches(1.1), 60, True, WHITE)
    add_text(
        s,
        "Your spending, explained in everyday words.",
        Inches(0.85),
        Inches(3.45),
        Inches(10),
        Inches(0.6),
        28,
        False,
        RGBColor(0xCB, 0xD5, 0xE1),
    )
    add_text(
        s,
        "A practice demo with sample bank data. No real bank login. No real card charge.",
        Inches(0.85),
        Inches(6.4),
        Inches(10),
        Inches(0.4),
        16,
        False,
        RGBColor(0x94, 0xA3, 0xB8),
    )
    notes(
        s,
        "Open with one sentence: Spendsense reads recent transactions and explains them the way a friend would, not the way a spreadsheet does.",
    )

    # 2 Problem
    s = new_slide(prs)
    header(s, "The everyday problem", "A bank app shows the numbers.\nIt does not explain them.")
    bullets(
        s,
        [
            "You can see that you spent more. You still have to hunt for why.",
            "A single big charge can hide in a long list.",
            "Next month's rent and salary are predictable, but the app rarely puts them on one picture.",
            "A chart does not tell you the one or two things worth changing.",
        ],
        top=Inches(2.15),
    )
    notes(s, "Stay on the feeling, not the technology. People already have statements. They do not have an explanation.")

    # 3 What it is
    s = new_slide(prs)
    header(s, "What it is", "A helper that sits on top of your transactions.")
    add_text(
        s,
        "Spendsense takes about three months of purchases and turns them into answers, warnings, a 30-day look ahead, savings ideas, and a short monthly checkup.",
        Inches(0.7),
        Inches(1.85),
        Inches(12),
        Inches(1.3),
        24,
        False,
        SLATE,
    )
    ideas = [
        ("Look back", "See where the money went, in plain groups."),
        ("Ask", "Type a normal question and get an answer from your own purchases."),
        ("Look ahead", "See the next 30 days, with room for uncertainty."),
    ]
    for i, (title, body) in enumerate(ideas):
        left = Inches(0.7 + i * 4.1)
        card(s, left, Inches(3.6), Inches(3.85), Inches(2.7))
        rect(s, left, Inches(3.6), Inches(3.85), Inches(0.12), EMERALD)
        add_text(s, title, left + Inches(0.28), Inches(3.95), Inches(3.3), Inches(0.5), 24, True, INK)
        add_text(s, body, left + Inches(0.28), Inches(4.6), Inches(3.3), Inches(1.3), 18, False, SLATE)
    notes(s, "Three jobs only: look back, ask, look ahead. Everything else is one of those.")

    # 4 What you can do
    s = new_slide(prs)
    header(s, "What you can do", "Eight things, one by one.")
    actions = [
        ("1", "Connect a bank", "Or use the practice bank already loaded."),
        ("2", "Read your list", "Each purchase has a group and a confidence score."),
        ("3", "Ask a question", "For example, “What are my loan payments?”"),
        ("4", "Check odd charges", "Unusual amounts are marked so you can look twice."),
        ("5", "See 30 days ahead", "Income, bills, and a range around the expected balance."),
        ("6", "Pick a savings idea", "Small cuts in optional spending, not in rent or loans."),
        ("7", "Read the monthly note", "A score plus exactly three next steps."),
        ("8", "Stay on free, or go Pro", "Free includes 10 questions a day. Pro removes that limit."),
    ]
    for i, (num, title, body) in enumerate(actions):
        col = i % 2
        row = i // 2
        left = Inches(0.7 + col * 6.25)
        top = Inches(1.8 + row * 1.25)
        add_text(s, num, left, top, Inches(0.55), Inches(0.5), 22, True, EMERALD)
        add_text(s, title, left + Inches(0.6), top, Inches(5.2), Inches(0.4), 20, True, INK)
        add_text(s, body, left + Inches(0.6), top + Inches(0.38), Inches(5.2), Inches(0.55), 16, False, MUTED)
    notes(s, "This is the menu for the live demo. Walk the screens in this order.")

    # 5 Groups
    s = new_slide(prs)
    header(s, "Your list of purchases", "Each line is grouped, and the app admits when it is unsure.")
    groups = [
        "Food",
        "Groceries",
        "Travel",
        "Shopping",
        "Entertainment",
        "Utilities",
        "Loan payments",
        "Income",
        "Health",
        "Other",
    ]
    for i, name in enumerate(groups):
        col = i % 5
        row = i // 5
        left = Inches(0.7 + col * 2.45)
        top = Inches(2.0 + row * 1.15)
        card(s, left, top, Inches(2.25), Inches(0.9))
        add_text(s, name, left + Inches(0.15), top + Inches(0.22), Inches(1.95), Inches(0.5), 18, True, INK, PP_ALIGN.CENTER)
    add_text(
        s,
        "Rent shows up as Other, on purpose. Housing is not one of the groups, so the app does not force it into a wrong box. A low score, under 60%, is highlighted in amber.",
        Inches(0.7),
        Inches(4.6),
        Inches(12),
        Inches(1.4),
        22,
        False,
        SLATE,
    )
    notes(
        s,
        "On the demo, point at Urban Living Apts, category Other, 58 percent. That is the unsure case.",
    )

    # 6 Ask
    s = new_slide(prs)
    header(s, "Ask in normal English", "The answer has to come from your purchases.")
    card(s, Inches(0.7), Inches(1.9), Inches(5.8), Inches(4.5))
    add_text(s, "A question that fits", Inches(1.0), Inches(2.15), Inches(5.2), Inches(0.45), 20, True, EMERALD)
    add_text(
        s,
        "“List my loan payments.”",
        Inches(1.0),
        Inches(2.8),
        Inches(5.2),
        Inches(0.7),
        26,
        True,
        INK,
    )
    add_text(
        s,
        "You get a short answer, plus the actual payments it used. The numbers are added up from those lines. The sentence does not invent a shop that is not in the list.",
        Inches(1.0),
        Inches(3.7),
        Inches(5.1),
        Inches(2.2),
        18,
        False,
        SLATE,
    )
    card(s, Inches(6.8), Inches(1.9), Inches(5.8), Inches(4.5))
    add_text(s, "A question that does not fit", Inches(7.1), Inches(2.15), Inches(5.2), Inches(0.45), 20, True, RGBColor(0xBE, 0x12, 0x30))
    add_text(
        s,
        "“What is the capital of France?”",
        Inches(7.1),
        Inches(2.8),
        Inches(5.2),
        Inches(0.9),
        26,
        True,
        INK,
    )
    add_text(
        s,
        "It says it does not have enough of your data to answer. There is no source list. It will not fill the gap with general knowledge.",
        Inches(7.1),
        Inches(3.95),
        Inches(5.1),
        Inches(2.0),
        18,
        False,
        SLATE,
    )
    notes(
        s,
        "In the mock demo, click List my EMI payments for the good answer. Then ask the France question. The long food question can also refuse, because the practice search needs shared words such as Chipotle. Say that out loud so it does not look like a bug.",
    )

    # 7 Odd charges
    s = new_slide(prs)
    header(s, "Unusual charges", "A warning when one purchase is far above the usual.")
    add_text(
        s,
        "For each group, Spendsense looks at the typical amount and the usual spread. If one charge is more than two steps above that typical amount, it is marked.",
        Inches(0.7),
        Inches(1.85),
        Inches(12),
        Inches(1.2),
        22,
        False,
        SLATE,
    )
    examples = [
        ("Delta Air Lines", "₹2,480", "A travel bill far above the usual trip."),
        ("Amazon", "₹1,899", "A shopping bill far above the usual order."),
        ("Chipotle", "₹420", "One meal far above the usual food bill."),
    ]
    for i, (name, amount, why) in enumerate(examples):
        left = Inches(0.7 + i * 4.1)
        card(s, left, Inches(3.4), Inches(3.85), Inches(2.6))
        add_text(s, name, left + Inches(0.25), Inches(3.6), Inches(3.4), Inches(0.45), 20, True, INK)
        add_text(s, amount, left + Inches(0.25), Inches(4.15), Inches(3.4), Inches(0.55), 28, True, EMERALD)
        add_text(s, why, left + Inches(0.25), Inches(4.85), Inches(3.35), Inches(0.9), 16, False, SLATE)
    notes(
        s,
        "Hover an alert in the live app and read the sentence. Smaller charges can also be marked if they still sit well above that group’s average.",
    )

    # 8 Forecast
    s = new_slide(prs)
    header(s, "The next 30 days", "A picture of money in, money out, and a range around it.")
    bullets(
        s,
        [
            "Salary, rent, and the loan payment are treated as things that come back on a schedule.",
            "The shaded band is wider on day 30 than on day 1, because the further out you look, the less sure the picture is.",
            "The written paragraph only repeats those numbers. It does not invent a new balance.",
            "The starting balance in the demo is an assumed ₹3,500, not a live bank balance.",
        ],
        top=Inches(1.9),
    )
    notes(s, "Point at the green line and the shaded area. Mention the day the expected balance goes negative, 6 October 2026 on this sample.")

    # 9 Savings + report
    s = new_slide(prs)
    header(s, "Ideas and a monthly checkup", "Small cuts, plus one score and three actions.")
    card(s, Inches(0.7), Inches(1.9), Inches(5.8), Inches(4.6))
    add_text(s, "Savings ideas", Inches(1.0), Inches(2.15), Inches(5.2), Inches(0.45), 24, True, INK)
    add_text(
        s,
        "Looks only at food, entertainment, and shopping, and only when there are enough of those purchases to learn from.\n\nThe suggestion is a 15% trim of what you already spend, not a new lifestyle. Rent, loan payments, and utility bills are left alone.\n\nYou can accept an idea or dismiss it.",
        Inches(1.0),
        Inches(2.8),
        Inches(5.1),
        Inches(3.3),
        18,
        False,
        SLATE,
    )
    card(s, Inches(6.8), Inches(1.9), Inches(5.8), Inches(4.6))
    add_text(s, "Monthly checkup", Inches(7.1), Inches(2.15), Inches(5.2), Inches(0.45), 24, True, INK)
    add_text(
        s,
        "A score from 0 to 100, using how much you saved, how spend compares with income, how many odd charges there were, and whether spend went up or down.\n\nAlways exactly three next steps. The paragraph restates the figures. It does not make up new ones.\n\nSeptember is the last full month in this demo.",
        Inches(7.1),
        Inches(2.8),
        Inches(5.1),
        Inches(3.3),
        18,
        False,
        SLATE,
    )
    notes(s, "Show one goal card, then the September report. Do not accept or dismiss a goal until you are finished talking, because a dismissed idea stays dismissed.")

    # 10 Free vs Pro and limits
    s = new_slide(prs)
    header(s, "Free, Pro, and the limits", "What this app will and will not do.")
    card(s, Inches(0.7), Inches(1.85), Inches(5.8), Inches(2.3))
    add_text(s, "Free", Inches(1.0), Inches(2.05), Inches(5), Inches(0.4), 22, True, INK)
    add_text(s, "Ten questions a day. A question that the app refuses still counts.", Inches(1.0), Inches(2.6), Inches(5.1), Inches(1.1), 18, False, SLATE)
    card(s, Inches(6.8), Inches(1.85), Inches(5.8), Inches(2.3))
    add_text(s, "Pro", Inches(7.1), Inches(2.05), Inches(5), Inches(0.4), 22, True, INK)
    add_text(s, "No daily question limit. In the practice demo, Upgrade writes this locally. No card is charged.", Inches(7.1), Inches(2.6), Inches(5.1), Inches(1.1), 18, False, SLATE)
    add_text(
        s,
        "It does not move money, pay a bill, or give investment advice. The 11th free question simply asks you to upgrade.",
        Inches(0.7),
        Inches(4.5),
        Inches(12),
        Inches(1.4),
        22,
        False,
        SLATE,
    )
    notes(s, "Click Upgrade to Pro only at the end of the demo. It changes the shared practice login for everyone using it.")

    # 11 Demo today
    s = new_slide(prs)
    header(s, "The practice demo", "Sample data, so you can click every screen safely.")
    bullets(
        s,
        [
            "About 90 days of made-up purchases, shown in rupees.",
            "The bank is already connected: First Platypus Bank.",
            "Sign in with demo@spendsense.dev and the password demo-pass-123.",
            "Open the app once before you present. The database sleeps when nobody is using it, and the first click after that is slow.",
            "Leave Connect bank alone during the talk. The data is already loaded.",
        ],
        top=Inches(1.85),
    )
    notes(s, "If a page looks empty, wait a moment and refresh. Do not re-seed or reconnect the bank in the middle of the demo.")

    # 12 Close
    s = new_slide(prs, dark=True)
    rect(s, 0, 0, Inches(0.22), H, EMERALD)
    add_text(s, "IN ONE SENTENCE", Inches(0.85), Inches(2.15), Inches(11), Inches(0.35), 16, True, EMERALD)
    add_text(
        s,
        "Spendsense explains money you already spent, flags the surprises, and sketches the next 30 days.",
        Inches(0.85),
        Inches(2.7),
        Inches(11),
        Inches(1.8),
        36,
        True,
        WHITE,
    )
    add_text(
        s,
        "Thank you",
        Inches(0.85),
        Inches(5.4),
        Inches(6),
        Inches(0.5),
        22,
        False,
        RGBColor(0x94, 0xA3, 0xB8),
    )
    notes(s, "Stop there. Offer to open the live app and walk the eight screens.")

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
